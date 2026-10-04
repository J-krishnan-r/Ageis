"""Deterministic readers for the heterogeneous Aegis evidence package.

Image regions and the one scanned PDF page use deliberately reviewed
transcriptions in sidecars. Their method is labelled ``human_transcription``
so an operator never mistakes a transcription for OCR or live telemetry.
"""

from __future__ import annotations

from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
import hashlib
import json
import re


@dataclass(frozen=True)
class Segment:
    segment_id: str
    source_path: str
    media_type: str
    locator: str
    text: str
    extraction_method: str
    source_sha256: str
    metadata: dict[str, Any]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _segment_id(relative: str, locator: str, text: str) -> str:
    key = f"{relative}\0{locator}\0{text}".encode("utf-8")
    return hashlib.sha256(key).hexdigest()[:20]


def _make_segment(
    path: Path,
    project_root: Path,
    locator: str,
    text: str,
    method: str,
    metadata: dict[str, Any] | None = None,
) -> Segment:
    relative = path.resolve().relative_to(project_root.parent.resolve()).as_posix()
    source_kind = path.suffix.lower().lstrip(".")
    normalized = re.sub(r"\s+", " ", text).strip()
    return Segment(
        segment_id=_segment_id(relative, locator, normalized),
        source_path=relative,
        media_type=source_kind,
        locator=locator,
        text=normalized,
        extraction_method=method,
        source_sha256=_sha256(path),
        metadata=metadata or {},
    )


class _StructuredHTML(HTMLParser):
    """Extract headings, paragraphs, list items, and rows without page chrome."""

    BLOCK_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "tr"}
    SKIP_TAGS = {"script", "style", "noscript"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.skip_depth = 0
        self.block: str | None = None
        self.cells: list[str] = []
        self.items: list[tuple[str, str]] = []
        self.current_cell = ""

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in self.SKIP_TAGS:
            self.skip_depth += 1
        elif self.skip_depth == 0 and tag in self.BLOCK_TAGS:
            self.block = tag
            self.cells = []
        elif self.skip_depth == 0 and tag in {"td", "th"} and self.block == "tr":
            self.current_cell = ""

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in self.SKIP_TAGS and self.skip_depth:
            self.skip_depth -= 1
        elif tag in {"td", "th"} and self.block == "tr":
            self.cells.append(self.current_cell.strip())
            self.current_cell = ""
        elif tag == self.block and self.skip_depth == 0:
            content = " | ".join(self.cells) if tag == "tr" else "".join(self.cells)
            content = re.sub(r"\s+", " ", content).strip()
            if content:
                self.items.append((tag, content))
            self.block = None
            self.cells = []

    def handle_data(self, data: str) -> None:
        if self.skip_depth == 0 and self.block:
            if self.block == "tr":
                self.current_cell += data
            else:
                self.cells.append(data)


class _VisibleHTML(HTMLParser):
    """Retain the complete text stream, including banners and warning callouts."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.skip_depth = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in _StructuredHTML.SKIP_TAGS:
            self.skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in _StructuredHTML.SKIP_TAGS and self.skip_depth:
            self.skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if self.skip_depth == 0 and data.strip():
            self.parts.append(data.strip())


def _extract_pdf(path: Path, root: Path, visual_evidence: dict[str, Any]) -> tuple[list[Segment], str | None]:
    from pypdf import PdfReader

    relative = path.resolve().relative_to(root.parent.resolve()).as_posix()
    reader = PdfReader(path)
    output: list[Segment] = []
    warnings: list[str] = []
    for page_no, page in enumerate(reader.pages, 1):
        extracted = (page.extract_text(extraction_mode="layout") or "").strip()
        if extracted:
            # Keep pages intact for cross-paragraph/table context. FTS5 searches
            # these pages as evidence units, while curated claims quote the span.
            output.append(_make_segment(path, root, f"Page {page_no}", extracted, "pypdf"))
        transcripts = visual_evidence.get(relative, [])
        matching = [item for item in transcripts if item.get("page") == page_no]
        if matching:
            for item in matching:
                output.append(
                    _make_segment(
                        path,
                        root,
                        f"Page {page_no}; {item['region']}",
                        item["text"],
                        "human_transcription",
                        {"reviewed": item.get("reviewed", False), "bbox_normalized": item.get("bbox_normalized")},
                    )
                )
            continue
        if extracted:
            continue
        warnings.append(f"{relative}: page {page_no} has no text layer and no reviewed transcription.")
    return output, " ".join(warnings) if warnings else None


def extract_document(path: Path, project_root: Path, sidecars: dict[str, Any]) -> list[Segment]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path, project_root, sidecars["visual_evidence"])[0]

    if suffix in {".html", ".htm"}:
        parser = _StructuredHTML()
        source = path.read_text(encoding="utf-8", errors="replace")
        parser.feed(source)
        segments: list[Segment] = []
        complete = _VisibleHTML()
        complete.feed(source)
        full_text = " ".join(complete.parts)
        if full_text:
            segments.append(_make_segment(path, project_root, "HTML document, complete visible text in source order", full_text, "html.parser"))
        counts: dict[str, int] = {}
        for tag, text in parser.items:
            counts[tag] = counts.get(tag, 0) + 1
            segments.append(_make_segment(path, project_root, f"HTML <{tag}> #{counts[tag]}", text, "html.parser"))
        return segments

    if suffix == ".json":
        parsed = json.loads(path.read_text(encoding="utf-8"))
        segments = []

        def visit(value: Any, key_path: str = "$") -> None:
            if isinstance(value, dict):
                for key, child in value.items():
                    visit(child, f"{key_path}.{key}")
            elif isinstance(value, list):
                for i, child in enumerate(value):
                    visit(child, f"{key_path}[{i}]")
            else:
                segments.append(_make_segment(path, project_root, key_path, f"{key_path} = {json.dumps(value, ensure_ascii=False)}", "json"))

        visit(parsed)
        return segments

    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=False, read_only=True)
        segments = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                populated = [cell for cell in row if cell.value is not None]
                if not populated:
                    continue
                values = " | ".join(f"{cell.coordinate}={cell.value}" for cell in populated)
                locator = f"Sheet {sheet.title}, row {row[0].row}, cells {populated[0].coordinate}:{populated[-1].coordinate}"
                segments.append(_make_segment(path, project_root, locator, values, "openpyxl"))
        return segments

    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        segments = []
        for index, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text.strip():
                style = paragraph.style.name if paragraph.style else "paragraph"
                segments.append(_make_segment(path, project_root, f"Paragraph {index} ({style})", paragraph.text, "python-docx"))
        for table_index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                text = " | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells)
                segments.append(_make_segment(path, project_root, f"Table {table_index}, row {row_index}", text, "python-docx"))
        for section_index, section in enumerate(document.sections, 1):
            for kind, part in (("header", section.header), ("footer", section.footer)):
                text = " | ".join(p.text.strip() for p in part.paragraphs if p.text.strip())
                if text:
                    segments.append(_make_segment(path, project_root, f"Section {section_index} {kind}", text, "python-docx"))
        return segments

    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        segments = []
        for slide_no, slide in enumerate(presentation.slides, 1):
            for shape_no, shape in enumerate(slide.shapes, 1):
                texts = []
                if shape.has_text_frame and shape.text.strip():
                    texts.append(shape.text.strip())
                if shape.has_table:
                    for row in shape.table.rows:
                        texts.append(" | ".join(cell.text.strip() for cell in row.cells))
                if texts:
                    segments.append(
                        _make_segment(path, project_root, f"Slide {slide_no}, shape {shape_no}", "\n".join(texts), "python-pptx")
                    )
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    segments.append(_make_segment(path, project_root, f"Slide {slide_no}, speaker notes", notes, "python-pptx"))
        return segments

    if suffix == ".png":
        relative = path.resolve().relative_to(project_root.parent.resolve()).as_posix()
        transcripts = sidecars["visual_evidence"].get(relative, [])
        segments = []
        for item in transcripts:
            segments.append(
                _make_segment(
                    path,
                    project_root,
                    item["region"],
                    item["text"],
                    "human_transcription",
                    {"reviewed": item.get("reviewed", False), "bbox_normalized": item.get("bbox_normalized")},
                )
            )
        if not segments:
            try:
                from PIL import Image

                with Image.open(path) as image:
                    dimensions = f"{image.width}x{image.height} pixels"
            except Exception as exc:
                dimensions = f"image metadata unavailable: {exc}"
            segments.append(_make_segment(path, project_root, "Image metadata", dimensions, "pillow_metadata"))
        return segments

    return []


def extract_corpus(project_root: Path, dataset_root: Path, sidecars: dict[str, Any]) -> tuple[list[Segment], dict[str, Any]]:
    all_segments: list[Segment] = []
    documents: list[dict[str, Any]] = []
    for path in sorted(p for p in dataset_root.rglob("*") if p.is_file() and p.suffix.lower() in {".pdf", ".html", ".htm", ".json", ".xlsx", ".xlsm", ".docx", ".pptx", ".png"}):
        before = len(all_segments)
        warning = None
        if path.suffix.lower() == ".pdf":
            page_segments, warning = _extract_pdf(path, project_root, sidecars["visual_evidence"])
            all_segments.extend(page_segments)
        else:
            all_segments.extend(extract_document(path, project_root, sidecars))
        relative = path.resolve().relative_to(project_root.parent.resolve()).as_posix()
        methods = sorted({s.extraction_method for s in all_segments[before:]})
        documents.append({
            "source_path": relative,
            "sha256": _sha256(path),
            "media_type": path.suffix.lower().lstrip("."),
            "segment_count": len(all_segments) - before,
            "extraction_methods": methods,
            "status": "indexed" if len(all_segments) > before else "metadata_only",
            "warning": warning,
        })
    report = {
        "dataset_file_count": len(documents),
        "segment_count": len(all_segments),
        "documents": documents,
        "warnings": [d["warning"] for d in documents if d["warning"]],
    }
    return all_segments, report
